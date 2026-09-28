"""Editable, self-contained Excel reports. Never mutate the open document.

Visible A:N follows the planilla; hidden O:AA holds contribution bookkeeping and
historical steel snapshots. Totals sum O, so FE members are never counted twice.
Formula caches come from the application's engine, not rounded display strings.
"""
from datetime import datetime
import math
import os
from pathlib import Path
import tempfile
import textwrap
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter as letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.page import PageMargins
from openpyxl.workbook.properties import CalcProperties

from metrado.hierarchy import materialize
from metrado import references as refs
from metrado.sheet import calculate, RESULT_COLUMN
from metrado.steel import description_base
from metrado.steel_config import DIAMETERS, dimensions
from metrado.swelling import volume_blocks, upgrade_legacy

PENDING = 'Pendiente'
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NUMBER_FORMAT = '#,##0.00;[Red]-#,##0.00;0.00'
INPUT_FORMAT = '0.00####;[Red]-0.00####;0.00'
COUNT_FORMAT = 'General'
FONT = 'Arial Narrow'
BODY_SIZE = 8
FIRST_ROW = 12
WIDTHS = (9.5, 34, 2.8, 3, 4.5, 4.5, 4.5, 3.2, 5.4, 5.4, 5.4, 5.4, 4.5, 7.5)
SUMMARY_WIDTHS = (12, 67, 6, 12)
LABELS = ('ÍTEM', 'DESCRIPCIÓN', 'Und.', 'Elem.', 'Largo',
          'Ancho', 'Alto', 'N.º de veces', 'Lon.',
          'Área', 'Vol.', 'Kg.', 'Und.', 'Total')
STEEL_ITEM_LABELS = {6: 'gancho', 7: 'empal.', 10: 'diam.', 11: 'kg/m'}


def _number(text, default=None):
    if text is None or not str(text).strip():
        return default
    try:
        value = float(str(text).strip().replace(',', '.'))
        return value if math.isfinite(value) else str(text)
    except (ValueError, OverflowError):
        return str(text)


def _literal(cell, value):
    if isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
        raise ValueError('Hay un texto demasiado largo o con caracteres de control no admitidos por Excel. Revisa las descripciones.')
    cell.value = value
    if isinstance(value, str):
        # Descriptions and identifiers are never interpreted as Excel formulas.
        cell.data_type = 's'


def _guard(expression, required, conditions=()):
    checks = [f'COUNT({",".join(required)})={len(required)}', *conditions]
    return f'IFERROR(IF(AND({",".join(checks)}),{expression},"{PENDING}"),"{PENDING}")'


class ExcelReport:
    def __init__(self, title, rows, engine):
        self.rows = materialize(rows, immutable_catalogs=True)
        upgrade_legacy(self.rows)
        refs.check_cycles(self.rows)
        self.values, self.errors = calculate(self.rows, engine)
        self.engine = engine
        self.title = title or 'Planilla de metrados'
        self.book = Workbook()
        self.book.remove(self.book.active)
        self.dev = self.book.create_sheet('Desarrollo')
        self.summary = self.book.create_sheet('Resumen')
        self.book.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
        self.book.properties.creator = 'Metrados'
        self.book.properties.title = self.title
        self.caches = {'Desarrollo': {}, 'Resumen': {}}
        self.positions, self.footers = {}, {}
        self.blocks = dict(volume_blocks(self.rows))
        self.ends = {end: start for start, end in self.blocks.items()}
        self.members = {i for start, end in self.blocks.items() for i in range(start, end + 1)}
        self.sources = {row['id']: row for row in self.rows if row['kind'] == 'item' and row.get('id')}
        self.catalogs = {}
        self.catalog_cursor = 2
        self._styles()
        self._layout()
        self._development()
        self._summary()

    def _styles(self):
        self.font = Font(name=FONT, size=BODY_SIZE, color='000000')
        self.input_font = self.font
        self.bold_font = Font(name=FONT, size=BODY_SIZE, bold=True, color='000000')
        self.line = Side(style='thin', color='BFBFBF')
        self.border = Border(left=self.line, right=self.line, top=self.line, bottom=self.line)
        self.frame = Side(style='medium', color='000000')
        self.colors = {'chapter': ('FFFFFF', '990099'), 'subchapter': ('FFFFFF', 'FF0000'),
                       'item': ('FFFFFF', '000000'), 'detail_group': ('FFFFFF', '000000'),
                       'factor': ('FFFFFF', '000000'), 'header': ('D9D9D9', '000000')}
        self.styles = {}
        for name in ('detail', *self.colors):
            self.styles[name] = {}
            for alignment in ('left', 'right'):
                style = NamedStyle(name=f'{name}_{alignment}', font=self.font, border=self.border,
                                   alignment=Alignment(horizontal=alignment, vertical='center',
                                                       wrap_text=alignment == 'left', shrink_to_fit=alignment == 'right'),
                                   number_format=NUMBER_FORMAT)
                if name in self.colors:
                    fill, color = self.colors[name]
                    style.fill = PatternFill('solid', fgColor=fill)
                    style.font = Font(name=FONT, size=BODY_SIZE, bold=True, color=color)
                self.book.add_named_style(style)
                self.styles[name][alignment] = style.name

    def _row(self, sheet, r, values=(), style=None):
        for c in range(1, 15 if sheet == self.dev else 5):
            cell = sheet.cell(r, c)
            cell.style = self.styles[style or 'detail']['left' if c <= 2 else 'right']
        for c, value in enumerate(values, 1):
            _literal(sheet.cell(r, c), value)
        sheet.row_dimensions[r].height = 14

    def _text_height(self, sheet, r, text, width, minimum=14):
        # The condensed reference font fits more characters without reducing its
        # point size. Explicit heights also work for merged Excel headings.
        lines = sum(max(1, len(textwrap.wrap(line, max(8, int(width * 1.3)), break_long_words=True)))
                    for line in str(text).split('\n'))
        sheet.row_dimensions[r].height = min(409, max(minimum, lines * 11 + 3))

    def _formula(self, sheet, address, expression, cached):
        sheet[address] = '=' + expression
        self.caches[sheet.title][address] = cached

    def _input(self, r, c, value, *, count=False):
        cell = self.dev.cell(r, c)
        _literal(cell, _number(value))
        cell.font = self.input_font
        cell.number_format = COUNT_FORMAT if count else INPUT_FORMAT

    def _layout(self):
        for sheet, widths, heading in (
                (self.dev, WIDTHS, 'METRADOS'),
                (self.summary, SUMMARY_WIDTHS, 'RESUMEN DE METRADOS')):
            last = letter(len(widths))
            sheet.sheet_view.showGridLines = False
            sheet.sheet_view.zoomScale = 90
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.sheet_properties.outlinePr.summaryRight = False
            sheet.page_setup.orientation = 'portrait'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.page_margins = PageMargins(left=.3, right=.3, top=.3, bottom=.35, header=.1, footer=.15)
            sheet.print_options.horizontalCentered = True
            sheet.oddFooter.center.text = 'Página &P de &N'
            sheet.oddFooter.center.size = 8
            sheet.oddFooter.center.font = FONT
            for c, width in enumerate(widths, 1):
                sheet.column_dimensions[letter(c)].width = width
            sheet.merge_cells(f'A1:{last}1')
            _literal(sheet['A1'], heading)
            sheet['A1'].font = Font(name=FONT, size=11, bold=True, color='000000')
            sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
            sheet.row_dimensions[1].height = 20
            self._frame_box(sheet, 1, 1, len(widths))
            for r in (2, 8, 11):
                sheet.row_dimensions[r].height = 6
            self._project_header(sheet, widths)
            sheet.freeze_panes = f'C{FIRST_ROW}'
            sheet.print_title_rows = f'1:{FIRST_ROW - 1}'
        for start, end, label in ((1, 1, 'ÍTEM'), (2, 2, 'DESCRIPCIÓN'), (3, 3, 'Und.'),
                                  (4, 4, 'Elem.'), (5, 7, 'DIMENSIONES'),
                                  (8, 8, 'N.º de\nveces'), (9, 13, 'METRADO'), (14, 14, 'Total')):
            self.dev.merge_cells(start_row=9, start_column=start, end_row=9 if start != end else 10, end_column=end)
            _literal(self.dev.cell(9, start), label)
        for r in (9, 10):
            for c in range(1, 15):
                cell = self.dev.cell(r, c)
                cell.fill = PatternFill('solid', fgColor='D9D9D9')
                cell.font = self.bold_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = self.border
            self.dev.row_dimensions[r].height = 16
        for c in (*range(5, 8), *range(9, 14)):
            _literal(self.dev.cell(10, c), LABELS[c - 1])
            self.dev.cell(10, c).alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
        for c in (3, 4, 8):
            self.dev.cell(9, c).alignment = Alignment(horizontal='center', vertical='center', text_rotation=90, wrap_text=True)
        self._frame_box(self.dev, 9, 10, 14)
        for c, label in enumerate(('ÍTEM', 'DESCRIPCIÓN', 'Und.', 'Total'), 1):
            self.summary.merge_cells(start_row=9, start_column=c, end_row=10, end_column=c)
            cell = self.summary.cell(9, c)
            _literal(cell, label)
            cell.font, cell.fill = self.bold_font, PatternFill('solid', fgColor='D9D9D9')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.border
            self.summary.cell(10, c).fill = PatternFill('solid', fgColor='D9D9D9')
            self.summary.cell(10, c).border = self.border
        self.summary.row_dimensions[9].height = self.summary.row_dimensions[10].height = 16
        self._frame_box(self.summary, 9, 10, 4)
        for c in range(15, 28):
            self.dev.column_dimensions[letter(c)].hidden = True
        for c, name in ((15, 'Aporte al total (no editar)'), (16, 'Descripción base'), (17, 'N.º ganchos'),
                        (18, 'Gancho personalizado (m)'), (19, 'Empalme unitario (m)'), (20, 'Gancho de catálogo (m)'),
                        (24, 'Diámetro'), (25, 'kg/m'), (26, 'Gancho (m)'), (27, 'Empalme (m)')):
            _literal(self.dev.cell(1, c), name)

    def _frame_box(self, sheet, first, last, columns):
        for r in range(first, last + 1):
            for c in range(1, columns + 1):
                cell = sheet.cell(r, c)
                border = cell.border
                cell.border = Border(left=self.frame if c == 1 else border.left,
                                     right=self.frame if c == columns else border.right,
                                     top=self.frame if r == first else border.top,
                                     bottom=self.frame if r == last else border.bottom)

    def _project_header(self, sheet, widths):
        last = letter(len(widths))
        for r, label in enumerate(('Proyecto:', 'Propietario:', 'Fecha:', 'Especialidad:', 'Módulo:'), 3):
            for c in range(1, len(widths) + 1):
                cell = sheet.cell(r, c)
                cell.font = self.font
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            _literal(sheet.cell(r, 1), label)
            sheet.row_dimensions[r].height = 14
        sheet.merge_cells(f'B3:{last}3')
        _literal(sheet['B3'], self.title)
        sheet['B3'].font = self.bold_font
        self._text_height(sheet, 3, self.title, sum(widths[1:]), 30)
        if sheet == self.dev:
            for r in range(4, 8):
                sheet.merge_cells(f'B{r}:H{r}')
                sheet.merge_cells(f'I{r}:K{r}')
                sheet.merge_cells(f'L{r}:N{r}')
            _literal(sheet['I4'], 'Hecho por:')
            _literal(sheet['I5'], 'Revisado por:')
        else:
            for r in range(4, 8):
                sheet.merge_cells(f'C{r}:D{r}')
            _literal(sheet['C4'], 'Hecho por:')
            _literal(sheet['C5'], 'Revisado por:')
        # Unknown project metadata stays blank; never borrow the reference's data.
        sheet['B5'] = datetime.now().date()
        sheet['B5'].number_format = 'dd/mm/yyyy'
        self._frame_box(sheet, 3, 7, len(widths))

    def _catalog(self, row):
        catalog = row.get('steel_catalog') if 'steel_hooks' in row else None
        if catalog is None:
            # Legacy bars retain the Rust density-based weights and manual lengths.
            catalog = {d: {'weight': str(self.engine.ask_sheet_steel(1., 0., 0., d, 1., 1., 1.)[1]),
                           'hook': '0', 'lap': '0'} for d in DIAMETERS}
            d = row['cells'][7]
            if d and d not in catalog:
                try:
                    catalog[d] = {'weight': str(self.engine.ask_sheet_steel(1., 0., 0., d, 1., 1., 1.)[1]),
                                  'hook': '0', 'lap': '0'}
                except ValueError:
                    pass
        key = tuple((d, *(entry[f] for f in ('weight', 'hook', 'lap'))) for d, entry in catalog.items())
        if key not in self.catalogs:
            start = self.catalog_cursor
            for d, entry in catalog.items():
                for c, value in zip(range(24, 28), (d, *(_number(entry[f]) for f in ('weight', 'hook', 'lap')))):
                    _literal(self.dev.cell(self.catalog_cursor, c), value)
                    self.dev.cell(self.catalog_cursor, c).font = self.input_font
                self.catalog_cursor += 1
            self.catalogs[key] = start, self.catalog_cursor - 1
        return self.catalogs[key], catalog

    def _development(self):
        cursor = FIRST_ROW
        for i, row in enumerate(self.rows):
            self.positions[i] = cursor
            cursor += 1
            if i in self.ends:
                self.footers[i] = cursor
                cursor += 1
        self.last_row = max(FIRST_ROW - 1, cursor - 1)
        source_positions = {row['id']: self.positions[i] for i, row in enumerate(self.rows)
                            if row['kind'] == 'item' and row.get('id')}
        for r in range(FIRST_ROW, cursor):
            self.dev.cell(r, 15, 0)
        owner, item_end = None, {}
        for i, row in enumerate(self.rows):
            r = self.positions[i]
            kind, cells = row['kind'], row['cells']
            if kind not in ('detail', 'detail_group'):
                if owner is not None:
                    item_end[owner] = r - 1
                owner = i if kind == 'item' else None
            style = ('subchapter' if row['level'] else 'chapter') if kind == 'chapter' else kind if kind != 'detail' else None
            self._row(self.dev, r, cells[:3], style)
            self._text_height(self.dev, r, cells[1], WIDTHS[1] - 3)
            self.dev.cell(r, 1).number_format = '@'
            self.dev.cell(r, 3).alignment = Alignment(horizontal='center', vertical='center')
            if kind == 'item':
                if cells[2] == 'kg':
                    for c, label in STEEL_ITEM_LABELS.items():
                        _literal(self.dev.cell(r, c), label)
                        self.dev.cell(r, c).alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
            elif kind == 'detail':
                self._detail(i, r, row, source_positions)
                result_col = 'L' if cells[2] == 'kg' else letter(RESULT_COLUMN[cells[2]] + 1)
                if i not in self.members:
                    self._formula(self.dev, f'O{r}', f'{result_col}{r}', self.values.get((i, RESULT_COLUMN[cells[2]]), PENDING))
            if i in self.ends:
                self._factor(i)
        if owner is not None:
            item_end[owner] = cursor - 1
        for i, end in item_end.items():
            r = self.positions[i]
            if end == r:
                expression = '0'
            else:
                contribution = f'O{r + 1}:O{end}'
                expression = f'IFERROR(IF(COUNT({contribution})={end-r},SUM({contribution}),"{PENDING}"),"{PENDING}")'
            self._formula(self.dev, f'N{r}', expression, self.values.get((i, 13), PENDING))
            self.dev[f'N{r}'].font = self.bold_font
        if not self.rows:
            self._row(self.dev, FIRST_ROW, ('', 'Sin partidas'))
            self.last_row = FIRST_ROW
        self.dev.print_area = f'A1:N{self.last_row}'
        self._page_breaks(self.dev, FIRST_ROW, self.last_row)

    def _detail(self, i, r, row, source_positions):
        cells, unit = row['cells'], row['cells'][2]
        self._input(r, 4, cells[3], count=True)
        steel = unit == 'kg' and not row['direct'] and 'reference' not in row
        if steel:
            self._steel(i, r, row)
            return
        times = cells[9] if unit == 'kg' and 'reference' not in row else cells[7]
        self._input(r, 8, times, count=True)
        required, conditions, terms = [f'D{r}', f'H{r}'], [f'D{r}<>0', f'H{r}<>0'], [f'D{r}', f'H{r}']
        if 'reference' in row:
            ref = row['reference']
            source = self.sources.get(ref['source'])
            _literal(self.dev[f'B{r}'], refs.description(row, self.sources))
            pos = source_positions.get(ref['source'])
            self._formula(self.dev, f'E{r}', f'N{pos}' if pos else f'"{PENDING}"', self.values.get((i, 4), PENDING))
            required.append(f'E{r}'); terms.append(f'E{r}')
            if (source['cells'][2] if source else ref['source_unit']) != unit:
                self._input(r, 6, ref['factor'])
                required.append(f'F{r}'); conditions.append(f'F{r}>0'); terms.append(f'F{r}')
                label = refs.description(row, self.sources).rsplit(' · ', 1)[0]
                _literal(self.dev[f'P{r}'], label + ' · ' + ref['factor_name'] + ': ')
                self._formula(self.dev, f'B{r}', f'P{r}&TEXT(F{r},"General")', refs.description(row, self.sources))
                try:
                    refs.conversion(ref, source['cells'][2] if source else ref['source_unit'], unit)
                except ValueError:
                    self._formula(self.dev, f'E{r}', f'"{PENDING}"', PENDING)
            self._text_height(self.dev, r, refs.description(row, self.sources), WIDTHS[1] - 3)
        elif unit == 'm3' and cells[9].strip():
            for c in range(5, 8):
                self._input(r, c, cells[c - 1])
            self._input(r, 10, cells[9])
            required.append(f'J{r}'); conditions.extend([
                f'J{r}>0', f'COUNT(E{r}:G{r})=1', f'SUM(E{r}:G{r})>0'])
            terms.extend([f'J{r}', f'SUM(E{r}:G{r})'])
        elif row['direct'].strip():
            self._input(r, 5, row['direct'])
            required.append(f'E{r}'); conditions.append(f'E{r}>0'); terms.append(f'E{r}')
            _literal(self.dev[f'F{r}'], 'DIRECTO')
            self.dev[f'F{r}'].font = Font(name=FONT, size=8, italic=True, color='000000')
        else:
            for c in range(5, 5 + {'m': 1, 'm2': 2, 'm3': 3}.get(unit, 0)):
                self._input(r, c, cells[c - 1])
                address = f'{letter(c)}{r}'
                required.append(address); conditions.append(f'{address}>0'); terms.append(address)
        col = 'L' if unit == 'kg' else letter(RESULT_COLUMN[unit] + 1)
        expression = '*'.join(terms)
        if 'reference' not in row:
            conditions.append(f'{expression}<>0')
        self._formula(self.dev, f'{col}{r}', _guard(expression, required, conditions), self.values.get((i, RESULT_COLUMN[unit]), PENDING))

    def _steel(self, i, r, row):
        cells = row['cells']
        self._input(r, 5, cells[4])
        bars, times = _number(cells[10]), _number(cells[9], 1)
        effective = bars * times if isinstance(bars, (int, float)) and isinstance(times, (int, float)) else None
        self._input(r, 8, effective if effective is not None and math.isfinite(effective) else PENDING, count=True)
        _literal(self.dev[f'J{r}'], cells[7])
        self.dev[f'J{r}'].font = self.input_font
        (start, end), catalog = self._catalog(row)
        match = f'MATCH(J{r},$X${start}:$X${end},0)'
        weight = f'INDEX($Y${start}:$Y${end},{match})'
        entry = catalog.get(cells[7], {})
        self._formula(self.dev, f'K{r}', f'IFERROR({weight},"{PENDING}")', self.values.get((i, 11), _number(entry.get('weight', ''), PENDING)))
        self.dev[f'K{r}'].number_format = '0.00'
        validation = DataValidation(type='list', formula1=f'$X${start}:$X${end}', allow_blank=True)
        validation.errorTitle, validation.error = 'Diámetro', 'Elige un diámetro de la lista.'
        validation.showErrorMessage = True
        self.dev.add_data_validation(validation); validation.add(f'J{r}')
        if 'steel_hooks' in row:
            config = row['steel_hooks']
            self._input(r, 17, config['count'], count=True)
            self._input(r, 18, config['override'])
            self._formula(self.dev, f'S{r}', f'IFERROR(INDEX($AA${start}:$AA${end},{match}),"{PENDING}")', _number(entry.get('lap', ''), PENDING))
            self._formula(self.dev, f'T{r}', f'IFERROR(INDEX($Z${start}:$Z${end},{match}),"{PENDING}")', _number(entry.get('hook', ''), PENDING))
            try:
                hook, lap, _ = dimensions(row)
                hook, lap = float(hook), float(lap)
            except (ValueError, KeyError):
                hook = lap = PENDING
            chosen = f'IF(R{r}="",T{r},R{r})'
            hook_expression = _guard(f'Q{r}*{chosen}', [f'Q{r}'],
                                     [f'OR(Q{r}=0,Q{r}=1,Q{r}=2)', f'ISNUMBER({chosen})', f'{chosen}>0'])
            self._formula(self.dev, f'F{r}', hook_expression, hook)
            # ROUND removes binary noise at exact 9 m boundaries; no lap feeds back into itself.
            lap_expression = f'MAX(0,ROUNDUP(ROUND((E{r}+F{r})/9,12),0)-1)*S{r}'
            self._formula(self.dev, f'G{r}', _guard(lap_expression, [f'E{r}', f'F{r}', f'S{r}'],
                                                   [f'E{r}>0', f'F{r}>=0', f'S{r}>0']), lap)
        else:
            self._input(r, 6, cells[5] or '0')
            self._input(r, 7, cells[6] or '0')
        required = [f'{c}{r}' for c in 'DEFGH']
        conditions = [f'D{r}<>0', f'E{r}>0', f'F{r}>=0', f'G{r}>=0', f'H{r}<>0']
        expr = f'D{r}*(E{r}+F{r}+G{r})*H{r}'
        self._formula(self.dev, f'I{r}', _guard(expr, required, conditions), self.values.get((i, 8), PENDING))
        self._formula(self.dev, f'L{r}', _guard(f'I{r}*K{r}', [f'I{r}', f'K{r}'], [f'K{r}>0', f'I{r}<>0']), self.values.get((i, 12), PENDING))
        _literal(self.dev[f'P{r}'], description_base(cells[1]))
        desc = f'{description_base(cells[1])} {effective:g} Ø{cells[7]}' if effective is not None else cells[1]
        self._formula(self.dev, f'B{r}', f'P{r}&" "&TEXT(H{r},"General")&" Ø"&J{r}', desc)
        self._text_height(self.dev, r, desc, WIDTHS[1] - 3)

    def _factor(self, end):
        start, r = self.ends[end], self.footers[end]
        config = self.rows[start]['volume_factor']
        self._row(self.dev, r, (), 'factor')
        _literal(self.dev[f'B{r}'], 'Factor de esponjamiento' + (' · ' + config['note'] if config['note'] else ''))
        self._text_height(self.dev, r, self.dev[f'B{r}'].value, WIDTHS[1] - 3)
        _literal(self.dev[f'C{r}'], 'm3')
        _literal(self.dev[f'I{r}'], 'FE')
        self._input(r, 10, config['factor'])
        a, b = self.positions[start], self.positions[end]
        quantity = f'K{a}:K{b}'
        expr = f'IFERROR(IF(AND(COUNT({quantity})={b-a+1},ISNUMBER(J{r}),J{r}>=1),SUM({quantity})*J{r},"{PENDING}"),"{PENDING}")'
        cached = self.values.get((end, 14), PENDING)
        self._formula(self.dev, f'K{r}', expr, cached)
        self._formula(self.dev, f'O{r}', f'K{r}', cached)

    def _summary(self):
        r = FIRST_ROW
        for i, row in enumerate(self.rows):
            if row['kind'] not in ('chapter', 'item'):
                continue
            cells = row['cells']
            style = ('chapter' if row['level'] == 0 else 'subchapter') if row['kind'] == 'chapter' else None
            self._row(self.summary, r, cells[:3], style)
            self.summary.cell(r, 1).number_format = '@'
            self._text_height(self.summary, r, cells[1], SUMMARY_WIDTHS[1] - 3)
            if row['kind'] == 'item':
                self._formula(self.summary, f'D{r}', f"'Desarrollo'!N{self.positions[i]}", self.values.get((i, 13), PENDING))
                self.summary[f'B{r}'].hyperlink = f"#'Desarrollo'!B{self.positions[i]}"
                self.summary[f'D{r}'].font = self.bold_font
            r += 1
        if r == FIRST_ROW:
            self._row(self.summary, FIRST_ROW, ('', 'Sin partidas'))
            r += 1
        self.summary.print_area = f'A1:D{r-1}'
        self._page_breaks(self.summary, FIRST_ROW, r-1)

    def _page_breaks(self, sheet, start, end):
        # Conservative printable body heights at nominal scale. Keep headings with
        # the next row; Excel also splits over-height blocks automatically.
        budget = (842 - (sheet.page_margins.top + sheet.page_margins.bottom) * 72
                  - sum(sheet.row_dimensions[r].height or 14 for r in range(1, start)) - 12)
        used = 0
        for r in range(start, end + 1):
            height = sheet.row_dimensions[r].height or 14
            following = (sheet.row_dimensions[r+1].height or 14) if r < end else 0
            is_heading = bool(sheet.cell(r, 2).font.bold)
            if used and used + height + (following if is_heading else 0) > budget:
                sheet.row_breaks.append(Break(id=r-1))
                used = 0
            used += height


def _save_cached(report, destination):
    """Write standard OOXML cached results without depending on Excel installation.

Only XML produced by our own Workbook is parsed. Automatic full recalculation
remains enabled; these caches make previews/data-only readers useful immediately.
"""
    with tempfile.TemporaryFile() as raw:
        report.book.save(raw)
        raw.seek(0)
        with ZipFile(raw) as source, ZipFile(destination, 'w', ZIP_DEFLATED) as target:
            names = {f'xl/worksheets/sheet{i}.xml': report.caches[sheet.title]
                     for i, sheet in enumerate(report.book.worksheets, 1)}
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename in names:
                    tree = ET.fromstring(data)
                    for cell in tree.iter(f'{{{NS}}}c'):
                        address = cell.get('r')
                        if address not in names[entry.filename]:
                            continue
                        cached = names[entry.filename][address]
                        value = cell.find(f'{{{NS}}}v')
                        if value is None:
                            value = ET.SubElement(cell, f'{{{NS}}}v')
                        cell.set('t', 'str' if isinstance(cached, str) else 'n')
                        value.text = str(cached)
                    data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
                target.writestr(entry, data)


def export_workbook(path, title, rows, engine):
    """Atomic export; the destination is unchanged if construction or saving fails."""
    report = ExcelReport(title, rows, engine)
    path, temporary = Path(path), None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix='.metrado-excel-', suffix='.tmp', delete=False) as stream:
            temporary = stream.name
            _save_cached(report, stream)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        return len([r for r in report.rows if r['kind'] == 'item']), len(report.errors)
    finally:
        report.book.close()
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
