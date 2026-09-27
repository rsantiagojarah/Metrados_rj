"""Read-only budget structure import. Never evaluate formulas or import amounts."""
from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile

from metrado.hierarchy import MAX_LEVEL, MAX_ROWS
from metrado.identity import ensure_ids
from metrado.sheet import UNITS, new_row, validate_project

MAX_SOURCE_ROWS = 30_000
MAX_COLUMNS = 128
HEADERS = ({'item', 'codigo', 'cod', 'nroitem'},
           {'descripcion', 'partida', 'partidas', 'descripciondelapartida'},
           {'unidad', 'und', 'unid', 'unidades'})


def normalized(value):
    text = unicodedata.normalize('NFKD', str(value or '').casefold())
    return ''.join(c for c in text if c.isalnum() and not unicodedata.combining(c))


def unit_code(value):
    key = normalized(value)
    aliases = {'ml': 'm', 'm1': 'm', 'metro': 'm', 'metros': 'm',
               'un': 'und', 'unid': 'und', 'unidad': 'und', 'unidades': 'und',
               'kgs': 'kg', 'kilogramo': 'kg', 'kilogramos': 'kg',
               'global': 'glb', 'viaje': 'vje', 'viajes': 'vje',
               'dias': 'dia', 'meses': 'mes'}
    result = aliases.get(key, key)
    if result not in UNITS:
        raise ValueError(f'Unidad no admitida: «{value}». Unidades válidas: {", ".join(UNITS)}.')
    return result


def item_code(cell):
    value = cell.value
    if isinstance(value, bool):
        raise ValueError('El ítem debe ser un código, no un valor lógico.')
    if isinstance(value, (int, float)):
        if not isinstance(value, int) and not value.is_integer():
            raise ValueError('Guarda el ítem como texto en Excel para conservar los puntos y ceros de su código.')
        value = str(int(value))
        if re.fullmatch(r'0{1,64}', cell.number_format or ''):
            value = value.zfill(len(cell.number_format))
    else:
        value = str(value or '').strip()
    if not re.fullmatch(r'[0-9]+(?:\.[0-9]+)*', value) or len(value) > 128:
        raise ValueError('Ítem inválido. Usa códigos como 01, 01.01 o 01.01.01, guardados como texto.')
    return value


@dataclass
class ImportResult:
    sheet: str
    rows: list
    titles: int
    items: int


class ExcelBudget:
    def __init__(self, path):
        self.path = Path(path)
        self.book = None
        if self.path.suffix.lower() != '.xlsx':
            raise ValueError('Selecciona un archivo .xlsx. Para archivos .xls, guarda antes una copia como .xlsx en Excel.')
        if self.path.stat().st_size > 20_000_000:
            raise ValueError('El Excel supera el límite de 20 MB.')
        try:
            with ZipFile(self.path) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 64_000_000:
                    raise ValueError('El contenido descomprimido del Excel supera el límite de 64 MB.')
            # Lazy import: opening Metrados does not pay Excel's startup cost.
            from openpyxl import load_workbook
            self.book = load_workbook(self.path, read_only=True, data_only=False, keep_links=False)
            self.headers = {}
            for sheet in self.book.worksheets:
                sheet.reset_dimensions()  # Do not trust stale producer dimensions.
                for number, row in enumerate(sheet.iter_rows(max_row=100, max_col=MAX_COLUMNS), 1):
                    labels = [normalized(cell.value) for cell in row]
                    positions = [[i for i, label in enumerate(labels) if label in names] for names in HEADERS]
                    if all(len(matches) == 1 for matches in positions):
                        self.headers[sheet.title] = (number, tuple(matches[0] for matches in positions))
                        break
            if not self.headers:
                raise ValueError('No se encontró una hoja con encabezados Ítem, Descripción (o Partida) y Unidad en las primeras 100 filas.')
        except Exception as error:
            self.close()
            if isinstance(error, (OSError, ValueError)):
                raise
            raise ValueError(f'No se pudo leer el archivo Excel: {error}') from error

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.close()

    def close(self):
        if self.book is not None:
            self.book.close()

    def read(self, sheet_name):
        header, columns = self.headers[sheet_name]
        rows, seen, stack = [], set(), []
        titles = items = 0
        sheet = self.book[sheet_name]
        try:
            for number, source in enumerate(sheet.iter_rows(min_row=header + 1,
                                                           max_col=max(columns) + 1), header + 1):
                if number > MAX_SOURCE_ROWS:
                    raise ValueError(f'La hoja supera el límite de {MAX_SOURCE_ROWS} filas de origen.')
                cells = [source[i] for i in columns]
                values = [cell.value for cell in cells]
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                    continue
                if all(normalized(v) in names for v, names in zip(values, HEADERS)):
                    continue  # Printed budgets sometimes repeat headers.
                try:
                    if any(cell.data_type in ('f', 'e') for cell in cells):
                        raise ValueError('Ítem, descripción y unidad deben contener valores, no fórmulas ni errores de Excel. Pega esas columnas como valores en una copia.')
                    code = item_code(cells[0])
                    description = str(values[1] if values[1] is not None else '').strip()
                    if not description:
                        raise ValueError('Falta la descripción.')
                    key = tuple(int(part) for part in code.split('.'))
                    if len(key) > MAX_LEVEL + 1:
                        raise ValueError('El ítem supera el límite de niveles de la planilla.')
                    if key in seen:
                        raise ValueError(f'El ítem {code} está duplicado.')
                    while stack and len(stack[-1][0]) >= len(key):
                        stack.pop()
                    if len(key) > 1:
                        if not stack or stack[-1][0] != key[:-1]:
                            raise ValueError(f'Falta el título padre de {code}, o está fuera de orden. Coloca cada título antes de sus hijos.')
                        if stack[-1][1] != 'chapter':
                            raise ValueError('Una partida con unidad no puede contener otras partidas. Su padre debe ser un título sin unidad.')
                    raw_unit = str(values[2] if values[2] is not None else '').strip()
                    kind = 'item' if raw_unit else 'chapter'
                    unit = unit_code(raw_unit) if raw_unit else ''
                    row = new_row(kind, code, description, unit, len(key) - 1)
                    rows.append(row)
                    seen.add(key)
                    stack.append((key, kind))
                    titles += kind == 'chapter'
                    items += kind == 'item'
                    if len(rows) > MAX_ROWS:
                        raise ValueError('La planilla admite hasta 10 000 filas importadas.')
                except ValueError as error:
                    raise ValueError(f'Hoja «{sheet_name}», fila {number}: {error}') from error
        except Exception as error:
            if isinstance(error, (OSError, ValueError)):
                raise
            raise ValueError(f'La hoja «{sheet_name}» está dañada o incompleta: {error}') from error
        if not items:
            raise ValueError(f'La hoja «{sheet_name}» no contiene partidas con unidad.')
        validate_project(dict(version=7, title=self.path.stem, rows=rows))
        ensure_ids(rows)
        return ImportResult(sheet_name, rows, titles, items)
