"""Area measurements accept any two of length, width and height."""
from itertools import combinations
import os
from pathlib import Path
import tempfile
import unittest

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from PySide6.QtCore import Qt
from openpyxl import load_workbook
import metrado._enlace as engine
from metrado.grid import SheetModel
from metrado.sheet import new_row, calculate, write_project, read_project
from metrado.excel_export import export_workbook


def area_rows():
    return [new_row('item', '01', 'MURO', 'm2', 0),
            new_row('detail', description='Cara', unit='m2', level=1)]


class AreaDimensionsTests(unittest.TestCase):
    def test_any_two_dimensions_with_factors_and_undo(self):
        for pair in combinations((4, 5, 6), 2):
            with self.subTest(pair=pair):
                model = SheetModel(engine, area_rows())
                for column in (4, 5, 6):
                    self.assertTrue(model.flags(model.index(1, column)) & Qt.ItemIsEditable)
                for column, value in zip(pair, ('2,5', '4')):
                    self.assertTrue(model.setData(model.index(1, column), value))
                model.setData(model.index(1, 3), '-2')
                model.setData(model.index(1, 7), '3')
                self.assertEqual(model.errors, {})
                self.assertEqual(model.values[(0, 13)], -60)
                model.setData(model.index(1, pair[1]), '')
                self.assertNotIn((0, 13), model.values)
                model.undo_stack.undo()
                self.assertEqual(model.values[(0, 13)], -60)

    def test_invalid_dimension_counts_and_values(self):
        for dimensions in (('', '', ''), ('2', '', ''), ('2', '3', '4'),
                           ('', '2', '0'), ('', '2', '-1'), ('', '2', 'nan'),
                           ('', '2', 'inf'), ('', '2', 'abc')):
            with self.subTest(dimensions=dimensions):
                rows = area_rows()
                rows[1]['cells'][4:7] = dimensions
                values, errors = calculate(rows, engine)
                self.assertIn(1, errors)
                self.assertNotIn((0, 13), values)

    def test_save_and_excel_keep_all_three_possible_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            for pair in combinations((4, 5, 6), 2):
                with self.subTest(pair=pair):
                    rows = area_rows()
                    for col, value in zip(pair, ('2.5', '4')):
                        rows[1]['cells'][col] = value
                    path = Path(folder) / 'area.json'
                    write_project(path, 'Áreas', rows)
                    restored = read_project(path)[1]
                    self.assertEqual(restored[1]['cells'][4:7], rows[1]['cells'][4:7])
                    self.assertEqual(calculate(restored, engine)[0][(0, 13)], 10)
                    xlsx = Path(folder) / 'area.xlsx'
                    export_workbook(xlsx, 'Áreas', rows, engine)
                    formulas = load_workbook(xlsx)
                    cached = load_workbook(xlsx, data_only=True)
                    try:
                        sheet = formulas['Desarrollo']
                        for col, value in zip(pair, (2.5, 4)):
                            self.assertEqual(sheet.cell(13, col + 1).value, value)
                        formula = sheet['J13'].value
                        self.assertIn('COUNT(E13:G13)=2', formula)
                        self.assertIn('COUNTA(E13:G13)=2', formula)
                        self.assertIn('COUNTIF(E13:G13,">0")=2', formula)
                        self.assertIn('PRODUCT(E13:G13)', formula)
                        self.assertEqual(cached['Desarrollo']['J13'].value, 10)
                    finally:
                        formulas.close()
                        cached.close()
