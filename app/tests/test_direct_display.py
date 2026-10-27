"""Individual direct measurements are visible without becoming dimensions."""
import os
from pathlib import Path
import tempfile
import unittest

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from PySide6.QtCore import Qt
from openpyxl import load_workbook
import metrado._enlace as engine
from metrado.grid import SheetModel
from metrado.sheet import new_row
from metrado.excel_export import NUMBER_FORMAT, export_workbook
from metrado.display import format_two_decimals


class DirectDisplayTests(unittest.TestCase):
    def model(self, unit):
        return SheetModel(engine, [new_row('item', '01', 'Prueba', unit, 0),
                                  new_row('detail', unit=unit, level=1)])

    def test_direct_label_value_readonly_factors_and_undo(self):
        for unit, label, result in [('m2', 'ÁREA', 9), ('m3', 'VOLUMEN', 10)]:
            with self.subTest(unit=unit):
                model = self.model(unit)
                model.setData(model.index(1, 3), '4')
                model.set_direct(1, '8,6275')
                self.assertEqual(model.data(model.index(1, 4)), label)
                self.assertEqual(model.data(model.index(1, 5)), '8,6275')
                self.assertFalse(model.flags(model.index(1, 5)) & Qt.ItemIsEditable)
                self.assertFalse(model.setData(model.index(1, 5), '20'))
                self.assertAlmostEqual(model.values[(1, result)], 34.51)
                model.setData(model.index(1, 7), '2')
                self.assertEqual(model.rows[1]['direct'], '8,6275')
                self.assertAlmostEqual(model.values[(1, result)], 69.02)
                model.set_direct(1, '10')
                self.assertEqual(model.data(model.index(1, 5)), '10')
                model.undo_stack.undo()
                self.assertEqual(model.data(model.index(1, 5)), '8,6275')
                model.set_direct(1, '')
                self.assertTrue(model.flags(model.index(1, 5)) & Qt.ItemIsEditable)
                self.assertEqual(model.data(model.index(1, 5)), '')

    def test_volume_from_area_stays_in_area_column(self):
        model = self.model('m3')
        model.set_direct(1, '10')
        self.assertTrue(model.setData(model.index(1, 9), '25'))
        self.assertTrue(model.setData(model.index(1, 5), '0,2'))
        self.assertEqual(model.rows[1]['direct'], '')
        self.assertEqual(model.data(model.index(1, 9)), '25.00')
        self.assertEqual(model.data(model.index(1, 5)), '0,2')
        self.assertEqual(model.values[(1, 10)], 5)

    def test_direct_length_uses_length_label_and_keeps_its_result_column(self):
        model = self.model('m')
        model.setData(model.index(1, 3), '4')
        model.set_direct(1, '8,6275')
        self.assertEqual(model.data(model.index(1, 4)), 'LONGITUD')
        self.assertAlmostEqual(model.values[(1, 8)], 34.51)

    def test_excel_matches_label_value_and_result(self):
        with tempfile.TemporaryDirectory() as folder:
            for unit, label, result in [('m2', 'ÁREA', 'J13'), ('m3', 'VOLUMEN', 'K13')]:
                model = self.model(unit)
                model.setData(model.index(1, 3), '4')
                model.set_direct(1, '8,6275')
                path = Path(folder) / 'direct.xlsx'
                export_workbook(path, 'Directo', model.rows, engine)
                book = load_workbook(path)
                cached = load_workbook(path, data_only=True)
                try:
                    sheet = book['Desarrollo']
                    self.assertEqual(sheet['E13'].value, label)
                    self.assertEqual(sheet['F13'].value, 8.6275)
                    self.assertIn('D13*H13*F13', sheet[result].value)
                    self.assertAlmostEqual(cached['Desarrollo'][result].value, 34.51)
                finally:
                    book.close()
                    cached.close()

    def test_excel_uses_length_label_for_direct_linear_measurement(self):
        model = self.model('m')
        model.setData(model.index(1, 3), '4')
        model.set_direct(1, '8,6275')
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'longitud.xlsx'
            export_workbook(path, 'Longitud', model.rows, engine)
            book = load_workbook(path)
            try:
                sheet = book['Desarrollo']
                self.assertEqual(sheet['F13'].value, 'LONGITUD')
                self.assertEqual(sheet['E13'].value, 8.6275)
                self.assertIn('D13*H13*E13', sheet['I13'].value)
            finally:
                book.close()

    def test_planilla_y_reportes_comparten_el_redondeo_visible(self):
        model = self.model('m3')
        for column, value in ((4, '17.40'), (5, '13.15'), (6, '1.50')):
            self.assertTrue(model.setData(model.index(1, column), value))

        with tempfile.TemporaryDirectory() as folder:
            formula_path = Path(folder) / 'coherente-formulas.xlsx'
            values_path = Path(folder) / 'coherente-valores.xlsx'
            export_workbook(formula_path, 'Coherente', model.rows, engine)
            export_workbook(values_path, 'Coherente', model.rows, engine, formulas=False)
            formula = load_workbook(formula_path)
            cached = load_workbook(formula_path, data_only=True)
            values = load_workbook(values_path)
            try:
                visible = model.data(model.index(0, 13))
                self.assertEqual(visible, '343.22')
                self.assertEqual(formula['Desarrollo']['N12'].number_format, NUMBER_FORMAT)
                self.assertEqual(format_two_decimals(cached['Desarrollo']['N12'].value), visible)
                self.assertEqual(format_two_decimals(cached['Resumen']['D12'].value), visible)
                self.assertEqual(values['Desarrollo']['N12'].number_format, NUMBER_FORMAT)
                self.assertEqual(format_two_decimals(values['Desarrollo']['N12'].value), visible)
                self.assertEqual(format_two_decimals(values['Resumen']['D12'].value), visible)
            finally:
                formula.close()
                cached.close()
                values.close()
