"""Manual and persisted area-base × one-dimension volume calculations."""
from contextlib import closing
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from openpyxl import load_workbook
from PySide6.QtCore import Qt

import metrado._enlace as engine
from metrado.database import SCHEMA_VERSION, read_database, write_database
from metrado.clipboard import ROW_MIME, decode_rows, encode_rows
from metrado.excel_export import export_workbook
from metrado.grid import SheetModel
from metrado.identity import ensure_ids
from metrado.sheet import calculate, new_row, read_project, write_project


def area_rows():
    rows = [new_row('item', '01', 'EXCAVACIÓN', 'm3', 0),
            new_row('detail', description='VEREDA IRREGULAR', unit='m3', level=1)]
    ensure_ids(rows)
    return rows


class AreaBaseTests(unittest.TestCase):
    def test_area_column_is_editable_and_any_one_dimension_completes_volume(self):
        for column in (4, 5, 6):
            with self.subTest(column=column):
                model = SheetModel(engine, area_rows())
                self.assertTrue(model.flags(model.index(1, 9)) & Qt.ItemIsEditable)
                self.assertTrue(model.setData(model.index(1, 9), '25,5'))
                self.assertEqual(model.data(model.index(1, 9)), '25.50')
                self.assertIn(1, model.errors)
                self.assertTrue(model.setData(model.index(1, column), '0,20'))
                self.assertAlmostEqual(model.values[(1, 10)], 5.1)
                self.assertAlmostEqual(model.values[(0, 13)], 5.1)

    def test_requires_exactly_one_positive_complementary_dimension(self):
        rows = area_rows()
        rows[1]['cells'][9] = '30'
        values, errors = calculate(rows, engine)
        self.assertEqual(values[(1, 9)], 30)
        self.assertIn('exactamente una dimensión', errors[1])

        rows[1]['cells'][4:6] = ['2', '3']
        values, errors = calculate(rows, engine)
        self.assertNotIn((1, 10), values)
        self.assertIn('exactamente una dimensión', errors[1])

        rows[1]['cells'][5] = ''
        values, errors = calculate(rows, engine)
        self.assertEqual(errors, {})
        self.assertEqual(values[(1, 10)], 60)

    def test_manual_area_edit_and_direct_quantity_are_independently_undoable(self):
        model = SheetModel(engine, area_rows())
        self.assertTrue(model.setData(model.index(1, 9), '12'))
        self.assertEqual(model.rows[1]['cells'][9], '12')
        model.undo_stack.undo()
        self.assertEqual(model.rows[1]['cells'][9], '')
        model.undo_stack.redo()
        model.set_direct(1, '7')
        self.assertEqual(model.rows[1]['cells'][9], '')
        self.assertEqual(model.rows[1]['direct'], '7')
        model.undo_stack.undo()
        self.assertEqual(model.rows[1]['cells'][9], '12')

    def test_row_clipboard_preserves_area_input(self):
        rows = area_rows()
        rows[1]['cells'][4] = '3'
        rows[1]['cells'][9] = '12.5'
        model = SheetModel(engine, rows)
        copied = decode_rows(encode_rows(model, [1]).data(ROW_MIME))
        self.assertEqual(copied[0]['cells'][9], '12.5')
        self.assertEqual(calculate([rows[0], copied[0]], engine)[0][(1, 10)], 37.5)

    def test_json_sqlite_and_excel_preserve_area_input_and_formula(self):
        rows = area_rows()
        rows[1]['cells'][5] = '0,10'
        rows[1]['cells'][9] = '125,50'
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            json_path = root / 'area.json'
            db_path = root / 'area.db'
            xlsx_path = root / 'area.xlsx'

            write_project(json_path, 'Áreas', rows)
            self.assertEqual(json.loads(json_path.read_text(encoding='utf-8'))['version'], 8)
            restored_json = read_project(json_path)[1]
            self.assertEqual(restored_json[1]['cells'][9], '125,50')

            write_database(db_path, 'Áreas', rows)
            restored_db = read_database(db_path)[1]
            self.assertEqual(restored_db[1]['cells'][9], '125,50')
            with closing(sqlite3.connect(db_path)) as connection:
                self.assertEqual(connection.execute('PRAGMA user_version').fetchone()[0], SCHEMA_VERSION)

            export_workbook(xlsx_path, 'Áreas', rows, engine)
            formulas = load_workbook(xlsx_path)
            cached = load_workbook(xlsx_path, data_only=True)
            try:
                self.assertEqual(formulas['Desarrollo']['J13'].value, 125.5)
                self.assertIn('COUNT(E13:G13)=1', formulas['Desarrollo']['K13'].value)
                self.assertAlmostEqual(cached['Desarrollo']['K13'].value, 12.55)
            finally:
                formulas.close()
                cached.close()


if __name__ == '__main__':
    unittest.main()
